#!/usr/bin/env python3
"""LIBER STATUS renderer — generates the house master-state workbook.
LAW: rendered view, not source of truth. Ledgers win. No hand-editing.
Reads SCHEMA.json (single source) + per-sheet CSV data files.
Emits: CSV (custody format) + xlsx (reading format)."""
import csv, json, hashlib, subprocess, sys
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
SCHEMA = json.loads((BASE / "SCHEMA.json").read_text())

def git_hash(path):
    try:
        return subprocess.run(["git", "-C", str(BASE), "rev-parse", "--short", "HEAD"],
                              capture_output=True, text=True).stdout.strip()
    except Exception:
        return "n/a"

def load_sheet(name):
    f = BASE / f"{name}.csv"
    if not f.exists():
        return SCHEMA["sheets"][name]["columns"], []
    with open(f, newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    cols = SCHEMA["sheets"][name]["columns"]
    return cols, rows

GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED   = PatternFill("solid", fgColor="FFC7CE")

def status_fill(cell, val):
    v = str(val).lower()
    if any(w in v for w in ("pass", "done", "closed", "complete", "active", "nominal", "ratified", "built", "flies", "green")):
        cell.fill = GREEN
    elif any(w in v for w in ("pending", "blocked", "gavel", "draft", "amber", "await", "provisional")):
        cell.fill = AMBER
    elif any(w in v for w in ("fail", "dead", "open", "red", "investigate", "stale", "orphan")):
        cell.fill = RED

def render():
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    src_hash = git_hash(BASE)
    wb = Workbook()
    wb.remove(wb.active)
    index_rows = []
    for name in SCHEMA["sheets"]:
        cols, rows = load_sheet(name)
        if name == "00_INDEX":
            continue
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [len(str(r[i-1])) for r in rows[:50]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
        if "status" in cols:
            si = cols.index("status") + 1
            for r in ws.iter_rows(min_row=2, min_col=si, max_col=si):
                status_fill(r[0], r[0].value or "")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        index_rows.append([name, f"{len(rows)} rows", SCHEMA['sheets'][name]['source'], src_hash])
    # 00_INDEX
    ws = wb.create_sheet("00_INDEX", 0)
    ws.append(["LIBER STATUS — house master state", "", "", ""])
    ws.append([f"generated_at: {now}  |  source_hash: {src_hash}  |  schema: {SCHEMA['schema_version']}", "", "", ""])
    ws.append(["LAW: rendered view, not source of truth. Ledgers win. No hand-editing.", "", "", ""])
    ws.append([])
    ws.append(["sheet", "rows", "source", "commit"])
    for c in ws[5]:
        c.font = Font(bold=True)
    for row in index_rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 16
    for L in "BCD":
        ws.column_dimensions[L].width = 22
    # CSV custody copies
    with open(BASE / "00_INDEX.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["generated_at", now]); w.writerow(["source_hash", src_hash])
        w.writerow(["sheet", "rows", "source", "commit"])
        w.writerows(index_rows)
    out = BASE / "SOULSMITH_MASTER_STATE.xlsx"
    wb.save(out)
    print(f"RENDER OK: {out} — {sum(len(load_sheet(n)[1]) for n in SCHEMA['sheets'] if n != '00_INDEX')} data rows, {now}")

if __name__ == "__main__":
    render()
